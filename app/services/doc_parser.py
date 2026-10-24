"""把用户上传的 JD / 简历文件提取成纯文本。

只依赖 pdfplumber 与 python-docx，两者都做懒加载，避免拖慢启动。
"""

from __future__ import annotations

import re
from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx"}

FILE_DIALOG_FILTER = (
    "支持的文件 (*.txt *.md *.pdf *.docx);;文本文件 (*.txt *.md);;PDF (*.pdf);;Word (*.docx)"
)


class ExtractError(Exception):
    pass


def _read_plain(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ExtractError("缺少 pdfplumber，无法解析 PDF") from exc

    chunks: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(text)
    return "\n".join(chunks)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractError("缺少 python-docx，无法解析 Word 文档") from exc

    document = docx.Document(str(path))
    chunks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u3000", " ").replace("\xa0", " ")
    # 去掉 PDF 常见的孤立页码行
    lines = []
    for line in text.split("\n"):
        stripped = line.strip()
        if re.fullmatch(r"[-—]?\s*\d{1,3}\s*[-—]?", stripped):
            continue
        lines.append(re.sub(r"[ \t]{2,}", " ", stripped))
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: str | Path) -> str:
    target = Path(path)
    if not target.exists():
        raise ExtractError(f"文件不存在：{target.name}")
    suffix = target.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ExtractError(f"暂不支持 {suffix or '该'} 格式，请转成 PDF、DOCX 或 TXT")

    if suffix in {".txt", ".md", ".markdown"}:
        raw = _read_plain(target)
    elif suffix == ".pdf":
        raw = _read_pdf(target)
    else:
        raw = _read_docx(target)

    cleaned = clean_text(raw)
    if len(cleaned) < 30:
        raise ExtractError("几乎没有提取到文字，可能是扫描件或图片型 PDF")
    return cleaned


def summarize_preview(text: str, limit: int = 220) -> str:
    flat = re.sub(r"\s+", " ", text).strip()
    return flat[:limit] + ("…" if len(flat) > limit else "")


_COMPANY_PATTERNS = [
    r"(?:公司|企业|单位)名称[：:]\s*([^\n，,。；;]{2,30})",
    r"^([\u4e00-\u9fa5A-Za-z0-9（）()]{2,20}(?:科技|集团|网络|信息|软件|数据|智能|银行|证券|医院|研究院)(?:有限公司|股份有限公司|公司)?)",
]

_POSITION_PATTERNS = [
    r"(?:职位|岗位|招聘职位|岗位名称)[：:]\s*([^\n，,。；;]{2,30})",
    r"(?:招聘|诚招|急招)\s*([\u4e00-\u9fa5A-Za-z]{2,20}(?:工程师|开发|专家|经理|主管|架构师|设计师|分析师|实习生))",
]


def guess_meta(jd_text: str) -> tuple[str, str]:
    """从 JD 里粗略猜公司和岗位，作为表单预填，用户可以改。"""
    company = ""
    position = ""
    for pattern in _COMPANY_PATTERNS:
        match = re.search(pattern, jd_text, re.MULTILINE)
        if match:
            company = match.group(1).strip()
            break
    for pattern in _POSITION_PATTERNS:
        match = re.search(pattern, jd_text, re.MULTILINE)
        if match:
            position = match.group(1).strip()
            break
    return company, position
